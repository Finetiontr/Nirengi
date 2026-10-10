"""Assemble the PowerPoint and PDF fallbacks of /sunum from capture.mjs layers.

    python build.py <layers dir> <out dir>

Every layer is a picture at its exact stage position, so a slide matches the web
deck. Blocks enter with Fade + a 16 px rise in their `--d` order (all automatic);
Niri and the bubble keep the names `!!niri` / `!!bubble` on every slide and the
slides change with Morph (Fade fallback), so Niri glides between marks. The film
slide carries the promo video (public/sunum), full-frame on its poster, and starts
it on its own as the slide opens. Speaker notes from notes.ts go into each slide's
notes pane, and the close also carries the short answers for likely questions.
Needs python-pptx and Pillow.
"""

import json
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

LAYERS = Path(sys.argv[1])
OUT = Path(sys.argv[2])
WORK = LAYERS / "_trim"
FILM = Path(__file__).resolve().parents[2] / "public" / "sunum" / "nirengi-film.mp4"

SLIDE_W, SLIDE_H = 12192000, 6858000  # 13.333 x 7.5 in
RISE = 16 / 900  # the web's translateY(16px) as a fraction of slide height
DUR = 450  # ms per block entrance (web: 520 ms)
MORPH = 500  # ms

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
P = "{%s}" % NS["p"]


def trim(src: Path, dst: Path, frame_w: int):
    """Crop a full-frame layer to its alpha box; returns (x, y, w, h) in EMU."""
    im = Image.open(src).convert("RGBA")
    box = im.getchannel("A").point(lambda a: 255 if a > 0 else 0).getbbox()
    if box is None:
        return None
    im.crop(box).save(dst, optimize=True)
    emu = SLIDE_W / frame_w
    x0, y0, x1, y1 = box
    return round(x0 * emu), round(y0 * emu), round((x1 - x0) * emu), round((y1 - y0) * emu)


class Timing:
    """Builds one slide's <p:timing>: everything plays on its own after the transition."""

    def __init__(self):
        self.next_id = 5  # 1 tmRoot, 2 mainSeq, 3 and 4 the click group that starts on its own
        self.effects = []
        self.media = None

    def nid(self):
        self.next_id += 1
        return self.next_id - 1

    def _set_visible(self, spid):
        return (
            f'<p:set><p:cBhvr><p:cTn id="{self.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="visible"/></p:to></p:set>'
        )

    def _fade(self, spid, dur):
        return (
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{self.nid()}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
        )

    def _anim(self, spid, attr, frm, to, dur, delay=0):
        return (
            f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base">'
            f'<p:cTn id="{self.nid()}" dur="{dur}" decel="100000" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav>'
            f'<p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>'
        )

    def _rot(self, spid, by, dur, delay):
        return (
            f'<p:animRot by="{by}"><p:cBhvr><p:cTn id="{self.nid()}" dur="{dur}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'
        )

    def _scale(self, spid, by, dur, delay, repeat=1000):
        return (
            f'<p:animScale><p:cBhvr><p:cTn id="{self.nid()}" dur="{dur}" autoRev="1" repeatCount="{repeat}" fill="hold">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr>'
            f'<p:by x="{by}" y="{by}"/></p:animScale>'
        )

    def _effect(self, preset, cls, delay, body):
        node = "afterEffect" if not self.effects else "withEffect"
        cid = self.nid()
        self.effects.append(
            f'<p:par><p:cTn id="{cid}" presetID="{preset}" presetClass="{cls}" presetSubtype="0" fill="hold" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body()}</p:childTnLst></p:cTn></p:par>'
        )

    def rise(self, spid, delay, dur=DUR):
        """Fade in while rising 16 px into place (the web's .s-in)."""
        self._effect(10, "entr", delay, lambda: self._set_visible(spid) + self._fade(spid, dur)
                     + self._anim(spid, "ppt_y", f"#ppt_y+{RISE:.4f}", "#ppt_y", dur))

    def pop_in(self, spid, delay, dur=380, start=0.7):
        """Fade in while growing from `start` to full size (a little overshoot feel via decel)."""
        self._effect(53, "entr", delay, lambda: self._set_visible(spid) + self._fade(spid, dur)
                     + self._anim(spid, "ppt_w", f"#ppt_w*{start}", "#ppt_w", dur)
                     + self._anim(spid, "ppt_h", f"#ppt_h*{start}", "#ppt_h", dur))

    def wave(self, spid, delay):
        """A quick teeter, Niri's hello."""
        steps = [(8, 110), (-16, 160), (14, 150), (-6, 120)]
        def body():
            out, t = "", 0
            for deg, ms in steps:
                out += self._rot(spid, deg * 60000, ms, t)
                t += ms
            return out
        self._effect(32, "emph", delay, body)

    def cheer(self, spid, delay):
        """Zirve: two grow/shrink pulses with a small teeter."""
        def body():
            out = self._scale(spid, 114000, 220, 0, repeat=2000)
            t = 0
            for deg, ms in [(7, 110), (-14, 160), (7, 110)]:
                out += self._rot(spid, deg * 60000, ms, t)
                t += ms
            return out
        self._effect(6, "emph", delay, body)

    def play(self, spid, ms):
        """Start a video as the slide opens (PowerPoint's Start: Automatically)."""
        self.media = spid
        self._effect(1, "mediacall", 0, lambda: (
            f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{self.nid()}" dur="{ms}" fill="hold"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>'))

    def xml(self):
        if not self.effects:
            return None
        media = "" if self.media is None else (
            f'<p:video><p:cMediaNode vol="80000"><p:cTn id="{self.nid()}" fill="hold" display="0">'
            '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{self.media}"/></p:tgtEl></p:cMediaNode></p:video>')
        return (
            f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(self.effects)
            + '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
            '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
            + media
            + '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
        )


def transition_xml(morph: bool):
    if not morph:
        return f'<p:transition xmlns:p="{NS["p"]}" spd="fast"><p:fade/></p:transition>'
    return (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        f'<p:transition xmlns:p="{NS["p"]}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="fast" p14:dur="{MORPH}">'
        '<p159:morph option="byObject"/></p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition xmlns:p="{NS["p"]}" spd="fast"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
    )


def main():
    manifest = json.loads((LAYERS / "manifest.json").read_text(encoding="utf-8"))
    frame_w = manifest["width"]
    WORK.mkdir(exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    placed = {"slides": []}
    close = max(x["n"] for x in manifest["slides"])

    for s in manifest["slides"]:
        n, src = s["n"], LAYERS / f"{s['n']:02d}"
        dst = WORK / f"{n:02d}"
        dst.mkdir(exist_ok=True)
        slide = prs.slides.add_slide(blank)
        shapes = slide.shapes
        rec = {"n": n, "pics": []}

        def put(file, name):
            box = trim(src / file, dst / file, frame_w)
            if box is None:
                return None
            pic = shapes.add_picture(str(dst / file), Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]))
            pic.name = name
            rec["pics"].append({"file": str(dst / file), "name": name, "box": box})
            return pic

        # The background is opaque and full-frame: no trim.
        bg = Image.open(src / "bg.png").convert("RGB")
        bg.save(dst / "bg.png", optimize=True)
        pic = shapes.add_picture(str(dst / "bg.png"), 0, 0, Emu(SLIDE_W), Emu(SLIDE_H))
        pic.name = f"s{n} zemin"
        rec["pics"].append({"file": str(dst / "bg.png"), "name": pic.name, "box": [0, 0, SLIDE_W, SLIDE_H]})

        t = Timing()
        blocks = []
        for layer in s["layers"]:
            p = put(layer["file"], f"s{n} blok {layer['k'] + 1}")
            if p is not None:
                blocks.append((layer["delay"], layer["k"], p.shape_id))
        if s.get("bare"):
            # The film: full frame on its poster, started by our own timing below.
            niri = bubble = None
            if FILM.exists():
                movie = shapes.add_movie(str(FILM), 0, 0, Emu(SLIDE_W), Emu(SLIDE_H), poster_frame_image=str(dst / "bg.png"), mime_type="video/mp4")
                movie.name = f"s{n} tanıtım videosu"
                old = slide._element.find(P + "timing")
                if old is not None:
                    old.getparent().remove(old)
                t.play(movie.shape_id, int(s.get("film_ms") or 90000))
        else:
            niri = put("niri.png", "!!niri")
            bubble = put("bubble.png", "!!bubble")

        # Blocks in their --d order; the first carries no wait after the transition.
        for delay, _, spid in sorted(blocks):
            t.rise(spid, int(delay))
        if n == 1:
            # Niri waves in on the opening, then the bubble pops.
            t.pop_in(niri.shape_id, 150)
            t.wave(niri.shape_id, 520)
            t.pop_in(bubble.shape_id, 420, dur=300, start=0.92)
        elif n == close:
            t.cheer(niri.shape_id, 250)

        sld = slide._element
        tr = etree.fromstring(transition_xml(morph=n > 1))
        timing_xml = t.xml()
        clr = sld.find(P + "clrMapOvr")
        anchor = clr if clr is not None else sld.find(P + "cSld")
        anchor.addnext(tr)
        if timing_xml:
            tr.addnext(etree.fromstring(timing_xml))

        notes = s.get("notes") or {}
        lines = [f"{n}. {s['title']} (~{notes.get('sec', '?')} sn)", f"Niri: {s['line']}", ""] + list(notes.get("say", []))
        if n == close:
            for a in manifest.get("qa", []):
                lines += ["", f"Soru: {a['q']}"] + list(a["say"])
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)
        placed["slides"].append(rec)

    pptx_path = OUT / "Nirengi-sunum.pptx"
    prs.save(pptx_path)
    (WORK / "placed.json").write_text(json.dumps(placed, indent=1), encoding="utf-8")

    # PDF: the full frames, final state.
    frames = [Image.open(LAYERS / f"{s['n']:02d}" / "full.png").convert("RGB") for s in manifest["slides"]]
    pdf_path = OUT / "Nirengi-sunum.pdf"
    frames[0].save(pdf_path, save_all=True, append_images=frames[1:], resolution=288.0, quality=92)
    print(f"{pptx_path} ({pptx_path.stat().st_size // 1024} KB), {pdf_path} ({pdf_path.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
